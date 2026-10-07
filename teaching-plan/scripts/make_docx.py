#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
make_docx.py —— 把教案 Markdown 转成排版精良的 Word 文档。

排版参照：kami 设计规范（色板/层次/留白节奏）→ 移植到 Word 样式。
  - 强调色 墨蓝 #1B365D（仅用于标题、色条、表头，占比很小）
  - 暖调灰阶 #3D3D3A / #504E49 / #6B6A64，不用冷灰
  - 表格用"三线表"（顶线/底线粗，表头下细线），学术出版通行样式
  - 标题左侧墨蓝竖条，正文宋体 11pt / 1.5 倍行距，A4 舒适页边距

用法：
    python3 make_docx.py input.md -o 教案-课题.docx --title "教案：XXX"
    python3 make_docx.py input.md --style kami     # 标题改用仓耳今楷/思源宋体

支持的 Markdown 子集：
    # 文档标题        ## 一级环节（H1）      ### 二级（H2）      #### 三级（H3）
    - / * 列表        > 提示卡               | 表格 |            ``` 代码块
    **行内加粗**       --- 分隔线（忽略）
"""
from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

# ---------------------------------------------------------------- 设计 token
BRAND = "1B365D"        # 墨蓝 · 唯一强调色
INK = "141413"          # 正文黑
CHARCOAL = "3D3D3A"     # 次级文字
OLIVE = "504E49"        # 说明文字
STONE = "6B6A64"        # 元信息
BORDER = "E8E6DC"       # 主分隔线
BORDER_SOFT = "E5E3D8"  # 次分隔线
TAG_BG = "EEF2F7"       # 表头底纹（墨蓝 8% 实色）
CARD_BG = "F7F6F1"      # 提示卡底纹（羊皮纸浅调）
CODE_BG = "F4F3EE"

# 字体：默认走中英文文档通行组合；--style kami 切换为 kami 衬线体系
FONT_PRESETS = {
    "standard": {"title": "黑体", "body": "宋体", "ascii": "Georgia"},
    "kami": {"title": "TsangerJinKai02", "body": "Source Han Serif SC", "ascii": "Charter"},
}


# ---------------------------------------------------------------- XML 小工具
def _set_run_font(run, ea: str, ascii_: str, size: float, color: str | None = None, bold: bool = False):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = ascii_
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), ascii_)
    rFonts.set(qn("w:hAnsi"), ascii_)
    rFonts.set(qn("w:eastAsia"), ea)


def _para_borders(p, edges: dict):
    """edges: {'bottom': {'sz':8,'color':BORDER,'space':4}}  单位 sz 为 1/8pt"""
    pPr = p._p.get_or_add_pPr()
    pBdr = pPr.get_or_add_pBdr() if hasattr(pPr, "get_or_add_pBdr") else None
    if pBdr is None:
        pBdr = OxmlElement("w:pBdr")
        pPr.append(pBdr)
    for edge in ("top", "left", "bottom", "right"):
        if edge not in edges:
            continue
        spec = edges[edge]
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), spec.get("val", "single"))
        e.set(qn("w:sz"), str(spec.get("sz", 8)))
        e.set(qn("w:space"), str(spec.get("space", 4)))
        e.set(qn("w:color"), spec.get("color", "auto"))
        pBdr.append(e)


def _para_shade(p, fill: str):
    pPr = p._p.get_or_add_pPr()
    shd = pPr.get_or_add_shd() if hasattr(pPr, "get_or_add_shd") else None
    if shd is None:
        shd = OxmlElement("w:shd")
        pPr.append(shd)
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)


def _cell_shade(cell, fill: str):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.get_or_add_shd() if hasattr(tcPr, "get_or_add_shd") else None
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)


def _cell_border(cell, edge: str, sz: int, color: str, val: str = "single"):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.get_or_add_tcBorders() if hasattr(tcPr, "get_or_add_tcBorders") else None
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcPr.append(borders)
    e = OxmlElement(f"w:{edge}")
    e.set(qn("w:val"), val)
    e.set(qn("w:sz"), str(sz))
    e.set(qn("w:space"), "0")
    e.set(qn("w:color"), color)
    borders.append(e)


def _table_borders(table, top_sz=14, bottom_sz=14, color=INK):
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")

    def edge(name, sz, val="single", col=color):
        e = OxmlElement(f"w:{name}")
        e.set(qn("w:val"), val)
        e.set(qn("w:sz"), str(sz))
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), col)
        return e

    borders.append(edge("top", top_sz))
    borders.append(edge("bottom", bottom_sz))
    borders.append(edge("left", 0, "none"))
    borders.append(edge("right", 0, "none"))
    borders.append(edge("insideH", 0, "none"))
    borders.append(edge("insideV", 0, "none"))
    tblPr.append(borders)


def _page_number(paragraph):
    run = paragraph.add_run()
    f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = "PAGE"
    f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "end")
    run._r.append(f1); run._r.append(it); run._r.append(f2)
    return run


# ---------------------------------------------------------------- Markdown 解析
def parse_markdown(text: str):
    """返回 block 列表: ('title'|'h1'|'h2'|'h3'|'p'|'li'|'quote'|'table'|'code', payload)"""
    blocks = []
    lines = text.replace("\r\n", "\n").split("\n")
    i = 0
    while i < len(lines):
        line = lines[i]
        s = line.strip()

        if s.startswith("```"):
            i += 1
            buf = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                buf.append(lines[i])
                i += 1
            i += 1
            blocks.append(("code", "\n".join(buf)))
            continue

        if s.startswith(":::slide"):
            label = s[len(":::slide"):].strip()
            i += 1
            buf = []
            while i < len(lines) and lines[i].strip() != ":::":
                buf.append(lines[i])
                i += 1
            i += 1
            blocks.append(("slide", (label, buf)))
            continue

        if s in ("\\newpage", "<!-- pagebreak -->", "<<<PAGE>>>"):
            blocks.append(("pagebreak", None))
            i += 1
            continue

        if s.startswith(":::cards"):
            parts = s.split()
            cols, row_h = 2, None
            for tok in parts[1:]:
                if tok.startswith("h="):
                    try:
                        row_h = float(tok[2:])
                    except ValueError:
                        pass
                elif tok.isdigit():
                    cols = max(1, int(tok))
            i += 1
            buf = []
            while i < len(lines) and lines[i].strip() != ":::":
                buf.append(lines[i])
                i += 1
            i += 1
            cards, cur = [], []
            for ln in buf:
                if ln.strip() in ("---", "===", "✂"):
                    cards.append(cur); cur = []
                else:
                    cur.append(ln)
            cards.append(cur)
            cards = [c for c in cards if any(x.strip() for x in c)]
            blocks.append(("cards", (cols, cards, row_h)))
            continue

        if s.startswith(":::gap"):
            try:
                pts = float(s.split()[1])
            except (IndexError, ValueError):
                pts = 12.0
            blocks.append(("gap", pts))
            i += 1
            continue

        if not s or set(s) <= {"-", "*", "_"} and len(s) >= 3:
            i += 1
            continue

        if s.startswith("|") and s.count("|") >= 2:
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                raw = lines[i].strip().strip("|")
                cells = [c.strip() for c in raw.split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c or "-") for c in cells):
                    rows.append(cells)
                i += 1
            if rows:
                blocks.append(("table", rows))
            continue

        m = re.match(r"^(#{1,4})\s+(.*)$", s)
        if m:
            level = len(m.group(1))
            key = {1: "title", 2: "h1", 3: "h2", 4: "h3"}[level]
            blocks.append((key, m.group(2).strip()))
            i += 1
            continue

        if s.startswith("> "):
            blocks.append(("quote", s[2:].strip()))
            i += 1
            continue

        if re.match(r"^[-*+]\s+", s):
            blocks.append(("li", re.sub(r"^[-*+]\s+", "", s)))
            i += 1
            continue

        if re.match(r"^\d+[.)]\s+", s):
            blocks.append(("li", re.sub(r"^\d+[.)]\s+", "", s)))
            i += 1
            continue

        blocks.append(("p", s))
        i += 1
    return blocks


_INLINE = re.compile(r"\*\*(.+?)\*\*")


def _add_inline(p, text: str, ea: str, ascii_: str, size: float, color: str, base_bold=False):
    pos = 0
    for m in _INLINE.finditer(text):
        if m.start() > pos:
            r = p.add_run(text[pos:m.start()])
            _set_run_font(r, ea, ascii_, size, color, base_bold)
        r = p.add_run(m.group(1))
        _set_run_font(r, ea, ascii_, size, color, True)
        pos = m.end()
    if pos < len(text):
        r = p.add_run(text[pos:])
        _set_run_font(r, ea, ascii_, size, color, base_bold)


_SLIDE_BG = "FFF6D8"      # 投屏块底纹（浅黄）
_SLIDE_TAG_BG = "FFF3C4"  # 投屏标签底纹
_SLIDE_BAR = "E3B505"     # 投屏左竖条（金）


def _render_slide(doc, payload, EA_BODY: str, ASCII_F: str):
    """投屏内容：黄底 + 金色左竖条 + 【投屏】标签；块内 --- 表示同组下一页。"""
    label, body = payload
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(9)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.2
    _para_shade(p, _SLIDE_TAG_BG)
    _para_borders(p, {"left": {"sz": 26, "color": _SLIDE_BAR, "space": 6}})
    r = p.add_run("▶ 投屏" + (f"｜{label}" if label else ""))
    _set_run_font(r, EA_BODY, ASCII_F, 8.5, "8A6D00", True)

    groups, cur = [], []
    for ln in body:
        if ln.strip() == "---":
            groups.append(cur)
            cur = []
        else:
            cur.append(ln)
    groups.append(cur)

    for gi, g in enumerate(groups):
        for ln in g:
            t = ln.rstrip()
            if not t.strip():
                continue
            p = doc.add_paragraph()
            pf = p.paragraph_format
            pf.space_before = Pt(0)
            pf.space_after = Pt(0)
            pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
            pf.line_spacing = 1.4
            pf.left_indent = Cm(0.3)
            _para_shade(p, _SLIDE_BG)
            _para_borders(p, {"left": {"sz": 26, "color": _SLIDE_BAR, "space": 6}})
            is_q = t.strip().startswith(("提问", "追问", "【提问】", "提问："))
            _add_inline(p, t, EA_BODY, ASCII_F, 10.5, BRAND if is_q else INK, base_bold=is_q)
        if gi < len(groups) - 1:
            sp = doc.add_paragraph()
            sp.paragraph_format.space_before = Pt(0)
            sp.paragraph_format.space_after = Pt(0)
            _para_shade(sp, _SLIDE_BG)
            _para_borders(sp, {"left": {"sz": 26, "color": _SLIDE_BAR, "space": 6}})
            rr = sp.add_run("⌁ 同组下一页投屏（PPT 换页）")
            _set_run_font(rr, EA_BODY, ASCII_F, 8, "B08A00")
    tail = doc.add_paragraph()
    tail.paragraph_format.space_before = Pt(0)
    tail.paragraph_format.space_after = Pt(4)


def _fitify(doc, scale: float = 0.90, margin_cm: float = 1.4):
    """材料包适配排版：字号略小、边距适中，让每份材料刚好落在一页内且看得清。"""
    for sec in doc.sections:
        sec.top_margin = sec.bottom_margin = Cm(margin_cm)
        sec.left_margin = sec.right_margin = Cm(margin_cm + 0.2)
        try:
            sec.header_distance = sec.footer_distance = Cm(0.8)
        except Exception:
            pass

    def tune(paras):
        for p in paras:
            pf = p.paragraph_format
            try:
                # 固定行距（:::gap 的空白段）不能被改成倍数行距
                if pf.line_spacing is not None and pf.line_spacing_rule != WD_LINE_SPACING.EXACTLY:
                    if float(pf.line_spacing) > 1.2:
                        pf.line_spacing = 1.32
                if pf.line_spacing_rule == WD_LINE_SPACING.EXACTLY:
                    pass  # 保留精确间距，不做压缩
                elif pf.space_after is not None and pf.space_after.pt > 6:
                    pf.space_after = Pt(4)
                if pf.space_before is not None and pf.space_before.pt > 8:
                    pf.space_before = Pt(6)
            except (TypeError, ValueError):
                pass
            for r in p.runs:
                if r.font.size:
                    r.font.size = Pt(max(8.0, round(r.font.size.pt * scale, 1)))

    tune(doc.paragraphs)
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                tune(c.paragraphs)
    for sec in doc.sections:
        tune(sec.header.paragraphs)
        tune(sec.footer.paragraphs)


def _compactify(doc, scale: float = 0.72, margin_cm: float = 1.0):
    """紧凑模式：压缩页边距、字号、段间距，把材料包控制在少量页内。"""
    for sec in doc.sections:
        sec.top_margin = sec.bottom_margin = Cm(margin_cm)
        sec.left_margin = sec.right_margin = Cm(margin_cm + 0.1)
        try:
            sec.header_distance = sec.footer_distance = Cm(0.6)
        except Exception:
            pass

    def shrink(paras):
        for p in paras:
            pf = p.paragraph_format
            try:
                if pf.line_spacing_rule == WD_LINE_SPACING.EXACTLY:
                    pass  # 保留 :::gap 的精确间距
                else:
                    if pf.space_after is not None and pf.space_after.pt > 3:
                        pf.space_after = Pt(pf.space_after.pt * 0.55)
                    if pf.space_before is not None and pf.space_before.pt > 3:
                        pf.space_before = Pt(pf.space_before.pt * 0.55)
                    if pf.line_spacing and float(pf.line_spacing) > 1.2:
                        pf.line_spacing = max(1.12, float(pf.line_spacing) * 0.85)
            except (TypeError, ValueError):
                pass
            for r in p.runs:
                if r.font.size:
                    r.font.size = Pt(max(7.0, round(r.font.size.pt * scale, 1)))

    shrink(doc.paragraphs)
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                shrink(c.paragraphs)
                # 压缩单元格内边距，行高能吃下不少空间
                tcPr = c._tc.get_or_add_tcPr()
                mar = OxmlElement("w:tcMar")
                for side, w in (("top", "8"), ("bottom", "8"), ("start", "60"), ("end", "60")):
                    e = OxmlElement(f"w:{side}")
                    e.set(qn("w:w"), w)
                    e.set(qn("w:type"), "dxa")
                    mar.append(e)
                tcPr.append(mar)
    for sec in doc.sections:
        shrink(sec.header.paragraphs)
        shrink(sec.footer.paragraphs)


def _table_grid(table, sz: int = 6, color: str = "BFBFBF"):
    """全网格细线：用于卡片布局，表框线就是裁切线。"""
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")

    def edge(name):
        e = OxmlElement(f"w:{name}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), str(sz))
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), color)
        return e

    for name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        borders.append(edge(name))
    tblPr.append(borders)


def _set_row_height(row, pt: float):
    trPr = row._tr.get_or_add_trPr()
    h = OxmlElement("w:trHeight")
    h.set(qn("w:val"), str(int(pt * 20)))  # twips
    h.set(qn("w:hRule"), "atLeast")
    trPr.append(h)


def _render_cards(doc, payload, EA_TITLE: str, EA_BODY: str, ASCII_F: str, big: bool = False):
    """多列卡片布局：卡片内容填进表格单元，边框即裁切线。payload=(列数, 卡片, 行高)"""
    cols, cards, row_h = payload
    rows = (len(cards) + cols - 1) // cols
    table = doc.add_table(rows=rows, cols=cols)
    table.autofit = True
    _table_grid(table)
    if row_h:
        for r in range(rows):
            _set_row_height(table.rows[r], row_h)
    title_pt = 12.5 if big else 10.5
    body_pt = 11.5 if big else 9.5
    for idx, card in enumerate(cards):
        r, c = divmod(idx, cols)
        cell = table.cell(r, c)
        cell.text = ""
        first = True
        for ln in card:
            t = ln.strip()
            if not t:
                continue
            p = cell.paragraphs[0] if first else cell.add_paragraph()
            first = False
            pf = p.paragraph_format
            pf.space_before = Pt(0)
            pf.space_after = Pt(4)
            pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
            pf.line_spacing = 1.3
            if t.startswith("###") or t.startswith("**"):
                txt = t.lstrip("#").strip().strip("*")
                _add_inline(p, txt, EA_TITLE, ASCII_F, title_pt, BRAND, base_bold=True)
            else:
                _add_inline(p, t, EA_BODY, ASCII_F, body_pt, INK)


# ---------------------------------------------------------------- 文档构建
def build(md_path: Path, out_path: Path, title: str | None, style: str = "standard",
          subtitle: str | None = None, meta_lines: list[str] | None = None,
          font_override: dict | None = None, compact: bool = False,
          header: bool = True, footer: bool = True, fit: bool = False):
    fonts = FONT_PRESETS.get(style, FONT_PRESETS["standard"])
    if font_override:
        fonts = {**fonts, **{k: v for k, v in font_override.items() if v}}
    EA_TITLE, EA_BODY, ASCII_F = fonts["title"], fonts["body"], fonts["ascii"]

    doc = Document()
    # 文档默认字体：保证未显式设字体的内容也落在同一套字里
    normal = doc.styles["Normal"]
    normal.font.name = ASCII_F
    normal.font.size = Pt(10.5)
    nrpr = normal.element.get_or_add_rPr()
    nrfonts = nrpr.find(qn("w:rFonts"))
    if nrfonts is None:
        nrfonts = OxmlElement("w:rFonts")
        nrpr.append(nrfonts)
    nrfonts.set(qn("w:ascii"), ASCII_F)
    nrfonts.set(qn("w:hAnsi"), ASCII_F)
    nrfonts.set(qn("w:eastAsia"), EA_BODY)
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    sec.top_margin, sec.bottom_margin = Cm(2.3), Cm(2.2)
    sec.left_margin, sec.right_margin = Cm(2.4), Cm(2.4)
    sec.header_distance, sec.footer_distance = Cm(1.2), Cm(1.2)

    blocks = parse_markdown(md_path.read_text(encoding="utf-8"))
    doc_title = title
    if not doc_title:
        for k, v in blocks:
            if k == "title":
                doc_title = v
                break
    doc_title = doc_title or md_path.stem

    # ---- 页眉：标题 + 细分隔线（灰）；header=False 时完全不写（发给学生的材料不要页眉）
    if header:
        hp = sec.header.paragraphs[0]
        hp.paragraph_format.space_after = Pt(2)
        r = hp.add_run(doc_title)
        _set_run_font(r, EA_BODY, ASCII_F, 8.5, STONE)
        _para_borders(hp, {"bottom": {"sz": 6, "color": BORDER, "space": 4}})

    # ---- 页脚：居中页码；footer=False 时不写
    if footer:
        fp = sec.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = _page_number(fp)
        _set_run_font(r, EA_BODY, ASCII_F, 9, STONE)

    first_title = True
    li_buffer: list[str] = []

    def flush_li():
        nonlocal li_buffer
        for item in li_buffer:
            p = doc.add_paragraph()
            pf = p.paragraph_format
            pf.left_indent = Cm(0.75)
            pf.first_line_indent = Cm(-0.4)
            pf.space_after = Pt(3)
            pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
            pf.line_spacing = 1.45
            _add_inline(p, "· " + item, EA_BODY, ASCII_F, 10.5, INK)
        li_buffer = []

    for kind, payload in blocks:
        if kind != "li":
            flush_li()

        if kind == "pagebreak":
            doc.add_page_break()
            continue

        if kind == "gap":
            p = doc.add_paragraph()
            pf = p.paragraph_format
            pf.space_before = Pt(0)
            pf.space_after = Pt(0)
            pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
            pf.line_spacing = Pt(float(payload))
            continue

        if kind == "cards":
            _render_cards(doc, payload, EA_TITLE, EA_BODY, ASCII_F, big=not compact)
            continue

        if kind == "slide":
            _render_slide(doc, payload, EA_BODY, ASCII_F)
            continue

        if kind == "title":
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6 if first_title else 12)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            _add_inline(p, payload, EA_TITLE, ASCII_F, 20, BRAND, base_bold=True)
            if first_title:
                first_title = False
                rule = doc.add_paragraph()
                rule.paragraph_format.space_after = Pt(14)
                _para_borders(rule, {"bottom": {"sz": 18, "color": BRAND, "space": 1}})
                if subtitle:
                    sp = doc.add_paragraph()
                    sp.paragraph_format.space_after = Pt(6)
                    _add_inline(sp, subtitle, EA_BODY, ASCII_F, 11, OLIVE)
                for line in (meta_lines or []):
                    mp = doc.add_paragraph()
                    mp.paragraph_format.space_after = Pt(2)
                    _add_inline(mp, line, EA_BODY, ASCII_F, 9.5, STONE)
            continue

        if kind in ("h1", "h2", "h3"):
            size = {"h1": 15, "h2": 12.5, "h3": 11.5}[kind]
            color = {"h1": BRAND, "h2": CHARCOAL, "h3": OLIVE}[kind]
            p = doc.add_paragraph()
            pf = p.paragraph_format
            pf.space_before = Pt(18 if kind == "h1" else 12)
            pf.space_after = Pt(6)
            pf.keep_with_next = True
            if kind == "h1":
                pf.left_indent = Cm(0.32)
                _para_borders(p, {"left": {"sz": 22, "color": BRAND, "space": 8}})
            _add_inline(p, payload, EA_TITLE, ASCII_F, size, color, base_bold=True)
            continue

        if kind == "quote":
            p = doc.add_paragraph()
            pf = p.paragraph_format
            pf.left_indent = Cm(0.5)
            pf.space_before = Pt(4)
            pf.space_after = Pt(6)
            pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
            pf.line_spacing = 1.45
            _para_borders(p, {"left": {"sz": 18, "color": BORDER_SOFT, "space": 8}})
            _para_shade(p, CARD_BG)
            _add_inline(p, payload, EA_BODY, ASCII_F, 10.5, OLIVE)
            continue

        if kind == "li":
            li_buffer.append(payload)
            continue

        if kind == "code":
            for cl in payload.split("\n"):
                p = doc.add_paragraph()
                pf = p.paragraph_format
                pf.left_indent = Cm(0.5)
                pf.space_after = Pt(0)
                _para_shade(p, CODE_BG)
                r = p.add_run(cl)
                _set_run_font(r, "Menlo", "Menlo", 9.5, CHARCOAL)
            continue

        if kind == "table":
            rows = payload
            ncol = max(len(r_) for r_ in rows)
            table = doc.add_table(rows=len(rows), cols=ncol)
            table.alignment = WD_TABLE_ALIGNMENT.LEFT
            table.autofit = True
            _table_borders(table)
            for ri, row in enumerate(rows):
                for ci in range(ncol):
                    cell = table.cell(ri, ci)
                    cell.text = ""
                    txt = row[ci] if ci < len(row) else ""
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_before = Pt(3)
                    p.paragraph_format.space_after = Pt(3)
                    p.paragraph_format.line_spacing = 1.3
                    if ri == 0:
                        _add_inline(p, txt, EA_BODY, ASCII_F, 10, BRAND, base_bold=True)
                        _cell_shade(cell, TAG_BG)
                        _cell_border(cell, "bottom", 8, BRAND)
                    else:
                        _add_inline(p, txt, EA_BODY, ASCII_F, 10, INK)
            doc.add_paragraph().paragraph_format.space_after = Pt(4)
            continue

        # 普通段落
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_after = Pt(6)
        pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        pf.line_spacing = 1.5
        _add_inline(p, payload, EA_BODY, ASCII_F, 11, INK)

    flush_li()
    if fit:
        _fitify(doc)
    elif compact:
        _compactify(doc)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return out_path, doc_title


def main():
    ap = argparse.ArgumentParser(description="教案 Markdown → 排版精良的 Word")
    ap.add_argument("input", help="输入 Markdown 文件")
    ap.add_argument("-o", "--out", help="输出 .docx 路径")
    ap.add_argument("--title", help="文档标题（缺省取第一个 # 标题）")
    ap.add_argument("--subtitle", help="标题下副信息")
    ap.add_argument("--meta", action="append", default=[], help="标题区元信息行，可重复")
    ap.add_argument("--style", choices=["standard", "kami"], default="standard",
                    help="字体体系：standard=黑体/宋体；kami=仓耳今楷/思源宋体")
    ap.add_argument("--font-title", help="覆盖标题中文字体名")
    ap.add_argument("--font-body", help="覆盖正文中文字体名")
    ap.add_argument("--font-ascii", help="覆盖西文字体名")
    ap.add_argument("--compact", action="store_true",
                    help="紧凑排版（缩小边距与字号），用于控制在少量页内")
    ap.add_argument("--fit", action="store_true",
                    help="适配排版：适中字号与边距，让每份材料刚好落在一页内（推荐用于材料包）")
    ap.add_argument("--no-header", action="store_true", help="不生成页眉（发给学生的材料用）")
    ap.add_argument("--no-footer", action="store_true", help="不生成页脚页码")
    a = ap.parse_args()

    src = Path(a.input)
    if not src.exists():
        print(f"输入文件不存在: {src}", file=sys.stderr)
        return 2
    out = Path(a.out) if a.out else src.with_suffix(".docx")
    path, doc_title = build(
        src, out, a.title, a.style, a.subtitle, a.meta,
        {"title": a.font_title, "body": a.font_body, "ascii": a.font_ascii},
        compact=a.compact, fit=a.fit,
        header=not a.no_header, footer=not a.no_footer,
    )
    print(f"已生成: {path}\n标题: {doc_title}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
