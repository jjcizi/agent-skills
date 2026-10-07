#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
make_pptx.py —— 从教案 Markdown 的 :::slide 投屏块生成投屏 PPT。

视觉参照 kami：羊皮纸底 #F5F4ED、唯一强调色墨蓝 #1B365D、衬线层次、大留白。
编排规则（对应 skill 的 §5.6）：
  - 一页只放一个核心问题；提问行渲染为墨蓝底白字的提问框
  - 块内用 `---` 分页 → 生成下一页，用于"逐条揭晓/逐轮加条件"
  - 第一行作为页面标题，其余为正文要点

用法：
    python3 make_pptx.py 教案.md -o 投屏PPT.pptx --title "教案：XXX"
"""
from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

# kami tokens
BG = "F5F4ED"        # 羊皮纸
BRAND = "1B365D"     # 墨蓝（唯一强调色）
INK = "141413"
CHARCOAL = "3D3D3A"
OLIVE = "504E49"
STONE = "6B6A64"
BORDER = "E8E6DC"
TAG = "B08A00"       # 页眉小标签（暖金）
QUESTION_BG = "1B365D"
QUESTION_FG = "FFFFFF"

FONT_PRESETS = {
    "standard": {"title": "黑体", "body": "宋体"},
    "kami": {"title": "TsangerJinKai02", "body": "Source Han Serif SC"},
}

W, H = Inches(13.333), Inches(7.5)   # 16:9
MARGIN = Inches(0.85)
CONTENT_W = W - MARGIN * 2


# ------------------------------------------------------------------ helpers
def _set_run(run, font: str, size: float, color: str, bold: bool = False):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = font
    run.font.color.rgb = RGBColor.from_string(color)
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set("typeface", font)


def _rect(slide, x, y, w, h, fill: str, line: bool = False, shape=MSO_SHAPE.RECTANGLE):
    sp = slide.shapes.add_shape(shape, int(x), int(y), int(w), int(h))
    sp.fill.solid()
    sp.fill.fore_color.rgb = RGBColor.from_string(fill)
    if line:
        sp.line.color.rgb = RGBColor.from_string(BORDER)
        sp.line.width = Pt(1)
    else:
        sp.line.fill.background()
    sp.shadow.inherit = False
    return sp


_INLINE = re.compile(r"\*\*(.+?)\*\*")


def _fill_para(p, text: str, font: str, size: float, color: str, bold: bool = False):
    p.alignment = PP_ALIGN.LEFT
    pos = 0
    for m in _INLINE.finditer(text):
        if m.start() > pos:
            r = p.add_run(); r.text = text[pos:m.start()]
            _set_run(r, font, size, color, bold)
        r = p.add_run(); r.text = m.group(1)
        _set_run(r, font, size, color, True)
        pos = m.end()
    if pos < len(text) or pos == 0:
        r = p.add_run(); r.text = text[pos:]
        _set_run(r, font, size, color, bold)


def _textbox(slide, x, y, w, h, text: str, font: str, size: float, color: str, bold=False):
    tb = slide.shapes.add_textbox(int(x), int(y), int(w), int(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    _fill_para(p, text, font, size, color, bold)
    return tb


# ------------------------------------------------------------------ parsing
def extract_slides(md_text: str):
    """返回 [(label, [page_lines, ...]), ...]；块内 --- 分页。"""
    lines = md_text.replace("\r\n", "\n").split("\n")
    out = []
    i = 0
    while i < len(lines):
        s = lines[i].strip()
        if s.startswith(":::slide"):
            label = s[len(":::slide"):].strip()
            i += 1
            body = []
            while i < len(lines) and lines[i].strip() != ":::":
                body.append(lines[i])
                i += 1
            i += 1
            pages, cur = [], []
            for ln in body:
                if ln.strip() == "---":
                    pages.append(cur); cur = []
                else:
                    cur.append(ln)
            pages.append(cur)
            out.append((label, [p for p in pages]))
        else:
            i += 1
    return out


# 投屏标签不得出现教学机制的"技法名"——那是教师视角，向学生会暴露设计意图
MECH_WORDS = {
    "悬念": "开场", "反转": "再看一次", "身份": "分组",
    "站队": "做选择", "盲盒": "抽签", "复盘": "回顾",
}


def sanitize_label(label: str):
    """把投屏标签中的机制名换成中性说法，返回 (新标签, 命中的机制词)。"""
    hits = []
    for k, v in MECH_WORDS.items():
        if k in label:
            label = label.replace(k, v)
            hits.append(k)
    return label, hits


def _is_question(t: str) -> bool:
    return t.strip().startswith(("提问", "追问", "【提问】"))


def _clean(t: str) -> str:
    t = t.rstrip()
    return re.sub(r"^[·•\-*]\s*", "", t.strip())


# ------------------------------------------------------------------ build
def build(md_path: Path, out_path: Path, title: str | None, style: str = "standard",
          subtitle: str | None = None):
    fonts = FONT_PRESETS.get(style, FONT_PRESETS["standard"])
    F_TITLE, F_BODY = fonts["title"], fonts["body"]

    slides = extract_slides(md_path.read_text(encoding="utf-8"))
    if not slides:
        print("警告：未找到 :::slide 投屏块", file=sys.stderr)

    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    blank = prs.slide_layouts[6]

    def new_slide():
        s = prs.slides.add_slide(blank)
        _rect(s, 0, 0, W, H, BG)
        # 顶部墨蓝细线
        _rect(s, 0, 0, W, Pt(6), BRAND)
        return s

    # ---- 封面
    if title:
        s = new_slide()
        _textbox(s, MARGIN, Inches(2.5), CONTENT_W, Inches(1.4), title, F_TITLE, 40, BRAND, True)
        rule = _rect(s, MARGIN, Inches(3.95), Inches(2.6), Pt(3), BRAND)
        if subtitle:
            _textbox(s, MARGIN, Inches(4.2), CONTENT_W, Inches(0.6), subtitle, F_BODY, 18, OLIVE)

    page_no = 0
    for label, pages in slides:
        label, hits = sanitize_label(label)
        if hits:
            print(f"⚠ 投屏标签含机制名 {hits}，已替换为中性说法"
                  f"（PPT 面向学生，不暴露教学策略）", file=sys.stderr)
        for plines in pages:
            body = [_clean(l) for l in plines if l.strip()]
            if not body and not label:
                continue
            page_no += 1
            s = new_slide()
            y = MARGIN

            if label:
                _textbox(s, MARGIN, y, CONTENT_W, Inches(0.4), label, F_BODY, 14, TAG, True)
                y += Inches(0.55)

            head = body[0] if body else ""
            rest = body[1:] if body else []
            if head and not _is_question(head):
                _textbox(s, MARGIN, y, CONTENT_W, Inches(1.1), head, F_TITLE, 30, BRAND, True)
                y += Inches(1.25)
            elif head:
                rest = body  # 整页就是一个提问

            # 字号自适应
            n = max(1, len(rest))
            size = 22 if n <= 4 else (20 if n <= 6 else (18 if n <= 8 else 16))
            gap = Inches(0.16)

            for t in rest:
                if _is_question(t):
                    h = Inches(0.95)
                    box = _rect(s, MARGIN, y, CONTENT_W, h, QUESTION_BG, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
                    tf = box.text_frame
                    tf.word_wrap = True
                    tf.margin_left = tf.margin_right = Inches(0.3)
                    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
                    p = tf.paragraphs[0]
                    _fill_para(p, t, F_BODY, 20, QUESTION_FG, True)
                    y += h + gap
                else:
                    tb = _textbox(s, MARGIN, y, CONTENT_W, Inches(0.6), t, F_BODY, size, INK)
                    # 估算高度：按文本长度折行
                    est_lines = max(1, int(len(t) / (26 if size >= 20 else 30)) + 1)
                    y += Inches(0.42) * est_lines + gap

            # 页脚
            _rect(s, MARGIN, H - Inches(0.75), CONTENT_W, Pt(1), BORDER)
            _textbox(s, MARGIN, H - Inches(0.6), Inches(6), Inches(0.35),
                     label or "", F_BODY, 11, STONE)
            _textbox(s, W - MARGIN - Inches(1.2), H - Inches(0.6), Inches(1.2), Inches(0.35),
                     f"{page_no:02d}", F_BODY, 11, STONE)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out_path))
    return out_path, len(prs.slides), len(slides)


def main():
    ap = argparse.ArgumentParser(description="教案 :::slide 块 → 投屏 PPT")
    ap.add_argument("input")
    ap.add_argument("-o", "--out")
    ap.add_argument("--title")
    ap.add_argument("--subtitle")
    ap.add_argument("--style", choices=["standard", "kami"], default="standard")
    a = ap.parse_args()
    src = Path(a.input)
    if not src.exists():
        print(f"输入不存在: {src}", file=sys.stderr)
        return 2
    out = Path(a.out) if a.out else src.with_suffix(".pptx")
    path, n_pages, n_blocks = build(src, out, a.title, a.style, a.subtitle)
    print(f"已生成: {path}\n投屏块 {n_blocks} 组 / 幻灯片 {n_pages} 页（含封面）")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
